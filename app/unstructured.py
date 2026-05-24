"""Text extraction helpers for unstructured invoice documents."""

from io import BytesIO
from pathlib import Path

SUPPORTED_UNSTRUCTURED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


def extract_text_from_document(file_name: str, content: bytes) -> str:
    """Extract text from a supported unstructured document upload."""

    suffix = Path(file_name or "").suffix.lower()
    if suffix not in SUPPORTED_UNSTRUCTURED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_UNSTRUCTURED_EXTENSIONS))
        raise ValueError(f"Unsupported file type. Use one of: {supported}.")

    if suffix in {".txt", ".md"}:
        return _decode_text(content)
    if suffix == ".pdf":
        return _extract_pdf_text(content)
    if suffix == ".docx":
        return _extract_docx_text(content)

    raise ValueError("Unsupported file type.")


def _decode_text(content: bytes) -> str:
    """Decode a text-like document using common encodings."""

    for encoding in ("utf-8-sig", "utf-8", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("Could not decode the text file.")


def _extract_pdf_text(content: bytes) -> str:
    """Extract embedded text from a PDF document."""

    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ValueError("Install pypdf to extract PDF text.") from exc

    reader = PdfReader(BytesIO(content))
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages).strip()
    if not text:
        raise ValueError("No embedded text found in PDF. OCR is required.")
    return text


def _extract_docx_text(content: bytes) -> str:
    """Extract text from a Word DOCX document."""

    try:
        from docx import Document
    except ImportError as exc:
        raise ValueError("Install python-docx to extract DOCX text.") from exc

    document = Document(BytesIO(content))
    chunks = [paragraph.text for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(cell.text for cell in row.cells))

    text = "\n".join(chunk for chunk in chunks if chunk.strip()).strip()
    if not text:
        raise ValueError("No text found in DOCX file.")
    return text
